"""STEP 3: Excel (.xlsx) customer import. Parses in memory only; never logs or stores file contents.

Validation reuses services.validate_customer_input (STEP 2 rules). Document tokens must match
a canonical key or label exactly (doc_types.strict_key); no fuzzy matching.
"""
import re

from dataclasses import dataclass, field
from io import BytesIO

from . import services
from .doc_types import strict_key

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 1000
DOC_SPLIT = re.compile(r"[,;|\n]+")

# canonical field -> accepted header spellings (compared after squashing case/space/underscore)
HEADERS: dict[str, tuple[str, ...]] = {
    "name": ("name", "full name"),
    "email": ("email", "email address"),
    "mobile": ("mobile", "mobile number"),
    "required_documents": ("required_documents", "required documents"),
    "send_consent": ("send_consent", "send consent email"),
}
REQUIRED_HEADERS = ("name", "email", "mobile", "required_documents")
TRUE_VALUES, FALSE_VALUES = {"", "yes", "y", "true", "1"}, {"no", "n", "false", "0"}


class ImportFileError(Exception):
    def __init__(self, code: str, message: str, status: int = 400):
        super().__init__(message)
        self.code, self.message, self.status = code, message, status


@dataclass
class RowResult:
    row: int
    status: str = "valid"  # valid | invalid | duplicate_customer | duplicate_excel_row
    errors: list[dict] = field(default_factory=list)
    name: str | None = None
    email: str | None = None
    mobile: str | None = None
    required_documents: list[str] = field(default_factory=list)
    send_consent: bool = True

    def out(self) -> dict:
        ok = self.status == "valid"
        return {
            "row": self.row, "status": self.status, "errors": self.errors,
            # normalized values only when the row is valid (avoid echoing malformed input)
            "name": self.name if ok else None, "email": self.email if ok else None,
            "mobile": self.mobile if ok else None,
            "required_documents": self.required_documents if ok else [],
            "send_consent": self.send_consent,
        }


def _squash(v) -> str:
    return re.sub(r"[\s_]+", "", str(v or "").strip().lower())


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)  # mobile numbers typed as numbers
    return str(v).strip()


def parse_workbook(filename: str | None, data: bytes) -> list[tuple[int, dict[str, str]]]:
    if not (filename or "").lower().endswith(".xlsx"):
        raise ImportFileError("invalid_file_type", "Upload an .xlsx file.")
    if len(data) > MAX_BYTES:
        raise ImportFileError("file_too_large", "File is larger than 5 MB.", 413)
    if not data:
        raise ImportFileError("empty_workbook", "The file is empty.")
    try:
        from openpyxl import load_workbook

        # read_only + data_only: cached values only, formulas/macros never evaluated or run
        wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception:  # corrupt zip, not a workbook, bad XML, ...
        raise ImportFileError("invalid_file", "The file could not be read as an Excel workbook.")
    try:
        ws = wb.worksheets[0] if wb.worksheets else None
        if ws is None:
            raise ImportFileError("empty_workbook", "The workbook has no sheets.")
        it = ws.iter_rows(values_only=True)
        header = next(it, None)
        if not header or all(_cell(h) == "" for h in header):
            raise ImportFileError("empty_workbook", "The workbook has no header row.")
        lookup = {_squash(a): f for f, aliases in HEADERS.items() for a in aliases}
        cols: dict[str, int] = {}
        for i, h in enumerate(header):
            f = lookup.get(_squash(h))
            if f and f not in cols:
                cols[f] = i
        missing = [h for h in REQUIRED_HEADERS if h not in cols]
        if missing:
            raise ImportFileError("missing_headers", f"Missing columns: {', '.join(missing)}.")
        rows: list[tuple[int, dict[str, str]]] = []
        for n, values in enumerate(it, start=2):
            rec = {f: _cell(values[i]) if i < len(values) else "" for f, i in cols.items()}
            if all(v == "" for v in rec.values()):
                continue  # blank row
            rows.append((n, rec))
            if len(rows) > MAX_ROWS:
                raise ImportFileError("too_many_rows", f"At most {MAX_ROWS} rows per file.")
        if not rows:
            raise ImportFileError("empty_workbook", "The workbook has no customer rows.")
        return rows
    finally:
        wb.close()


def evaluate(db, rows: list[tuple[int, dict[str, str]]]) -> list[RowResult]:
    results: list[RowResult] = []
    for n, rec in rows:
        r = RowResult(row=n)
        consent = rec.get("send_consent", "").lower()
        if consent in TRUE_VALUES:
            r.send_consent = True
        elif consent in FALSE_VALUES:
            r.send_consent = False
        else:
            r.errors.append({"code": "invalid_send_consent", "message": "Send consent must be Yes or No."})
        keys: list[str] = []
        for tok in (t.strip() for t in DOC_SPLIT.split(rec.get("required_documents", ""))):
            if not tok:
                continue
            k = strict_key(tok)
            if k is None:
                r.errors.append({"code": "unsupported_document_type", "message": f"Unsupported document type: {tok}"})
            elif k not in keys:
                keys.append(k)
        try:
            r.name, r.email, r.mobile, r.required_documents = services.validate_customer_input(
                db, rec.get("name", ""), rec.get("email", ""), rec.get("mobile") or None, keys)
        except services.CustomerError as e:
            r.errors.append({"code": e.code, "message": e.message})
        codes = {e["code"] for e in r.errors}
        if codes - {"duplicate_customer"}:
            r.status = "invalid"
        elif codes:
            r.status = "duplicate_customer"
        results.append(r)

    # in-file duplicates (by normalized email); every occurrence is flagged
    seen: dict[str, int] = {}
    for r in results:
        key = _email_key(r, rows)
        if key:
            seen[key] = seen.get(key, 0) + 1
    for r in results:
        key = _email_key(r, rows)
        if key and seen[key] > 1:
            r.errors.append({"code": "duplicate_excel_row", "message": "This email appears more than once in the file."})
            if r.status == "valid":
                r.status = "duplicate_excel_row"
    return results


def _email_key(r: RowResult, rows: list[tuple[int, dict[str, str]]]) -> str:
    if r.email:
        return r.email
    raw = next((rec.get("email", "") for n, rec in rows if n == r.row), "")
    return raw.strip().lower()


def summary(results: list[RowResult]) -> dict:
    return {
        "total_rows": len(results),
        "valid_rows": sum(r.status == "valid" for r in results),
        "invalid_rows": sum(r.status == "invalid" for r in results),
        "duplicate_rows": sum(r.status in ("duplicate_customer", "duplicate_excel_row") for r in results),
    }
